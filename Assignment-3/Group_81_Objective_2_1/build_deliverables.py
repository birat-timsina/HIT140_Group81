"""Create the Word report, feature dictionary, executed notebook and slide data."""
from pathlib import Path
import json
import subprocess
import sys
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
import contextlib
import io
import project as p

ROOT=Path(__file__).resolve().parent
OUT=ROOT/'deliverables'
OUT.mkdir(exist_ok=True)

FEATURE_INFO=[
 ('Competitive form','competitive_form_gap','Average result points from the previous ten selected competitive internationals. Award 3 for a win, 1 for a draw and 0 for a loss, regardless of competition format.',
  'Summarises recent competitive success without requiring a large winning margin. Higher values for A are expected to favour A.',
  'Points ignore the margin of victory and remain sensitive to opponent quality. A draw followed by a shootout win still earns one synthetic result point.'),
 ('Clean sheet rate','clean_sheet_gap','Fraction of the previous ten competitive matches in which the team conceded zero goals.',
  'Measures how often a defence completely prevents scoring. This can distinguish teams with similar average concession totals.',
  'It discards the difference between conceding one and several goals. Opponents and tactical choices affect the rate.'),
 ('Scoring failure rate','scoring_failure_gap','Fraction of the previous ten competitive matches in which the team scored zero goals.',
  'Measures attacking reliability. A higher failure rate for A is expected to lower its goal difference.',
  'A scoreless match can reflect poor finishing or strong opposition. This is an observed event rate, not shot-based expected goals.'),
 ('Margin volatility','margin_volatility_gap','Sample standard deviation of goals scored minus goals conceded over the previous ten competitive games, using n minus 1 in the variance denominator.',
  'Tests whether inconsistency adds information beyond form. Large wins and losses can distinguish unstable performance.',
  'There is no definite expected coefficient sign. Volatility may affect uncertainty more than the expected margin, so this is deliberately an exploratory feature.'),
 ('Form momentum','form_momentum_gap','Mean result points in the most recent five competitive games minus the mean in the preceding five. Then subtract the two teams\' momentum values.',
  'Tests whether a recent improvement adds information beyond the ten-match average.',
  'A short run can be random or reflect easier opponents. The feature shares information with competitive form; usefulness must be tested conditionally.'),
 ('Past opponent strength','schedule_strength_gap','For each of the ten previous competitive opponents, calculate its points per game in its own ten competitive games BEFORE that historical meeting. Average those ten opponent values.',
  'Adds context to the schedule behind a team\'s form. Similar results against stronger opponents may mean something different.',
  'This is a simple schedule proxy, not an Elo rating or a complete strength adjustment. Strong teams in weak confederations can still have high points rates.'),
 ('Recent workload','workload_30d_gap','Count all recorded internationals in the interval from 30 days before the prediction date up to, but excluding, that date. Include friendlies for this feature.',
  'Tests a scheduling and recovery hypothesis. More recent matches may mean more fatigue, but could also improve preparation.',
  'The count does not measure player minutes, club workload, travel or rotation. Fixture-congestion evidence is mixed, and a 30-day count is a coarse proxy (Julian et al., 2021).'),
 ('World Cup experience','world_cup_experience_gap','Count matches played in the 2014, 2018 and 2022 World Cups and apply the natural logarithm of one plus that count. Subtract B from A.',
  'Tests whether recent national-team tournament experience adds context. The logarithm limits the influence of highly experienced nations and handles zero appearances.',
  'This measures national-team history, not the current players\' experience. It also reflects prior qualification and strength; no causal interpretation is justified.'),
]

def create_report():
    summary=json.loads((ROOT/'results/summary.json').read_text())
    ab=pd.read_csv(ROOT/'results/feature_ablation.csv').set_index('feature')
    coeff=pd.read_csv(ROOT/'results/coefficients.csv')
    metrics=pd.read_csv(ROOT/'results/metrics.csv')
    sensitivity=pd.read_csv(ROOT/'results/sensitivity.csv')
    doc=Document();sec=doc.sections[0]
    sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=Inches(.7)
    sec.left_margin=sec.right_margin=Inches(.8)
    for sty in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
        st=doc.styles[sty];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0)
    normal=doc.styles['Normal'];normal.font.size=Pt(11)
    normal.paragraph_format.space_after=Pt(7);normal.paragraph_format.line_spacing=1.08
    doc.styles['Title'].font.size=Pt(27)
    doc.styles['Heading 1'].font.size=Pt(18);doc.styles['Heading 2'].font.size=Pt(13)
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    run=footer.add_run('Group 81  |  Objective 2.1  |  ');run.font.size=Pt(9)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    def para(t): return doc.add_paragraph(t)
    def heading(t): doc.add_heading(t,1)
    def new(t): doc.add_page_break();heading(t)
    def table(headers,rows,widths):
        tb=doc.add_table(rows=1,cols=len(headers));tb.autofit=False
        for i,h in enumerate(headers):tb.rows[0].cells[i].text=str(h)
        for row in rows:
            cells=tb.add_row().cells
            for i,val in enumerate(row):cells[i].text=str(val)
        for ri,row in enumerate(tb.rows):
            for ci,cell in enumerate(row.cells):
                cell.width=Inches(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                pr=cell._tc.get_or_add_tcPr();shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'173B48' if ri==0 else ('EEF3F5' if ri%2==0 else 'FFFFFF'));pr.append(shade)
                borders=OxmlElement('w:tcBorders')
                for side in ['top','left','bottom','right']:
                    edge=OxmlElement('w:'+side);edge.set(qn('w:val'),'single');edge.set(qn('w:sz'),'4');edge.set(qn('w:color'),'D9D9D9');borders.append(edge)
                pr.append(borders)
                margins=OxmlElement('w:tcMar')
                for side in ['top','left','bottom','right']:
                    edge=OxmlElement('w:'+side);edge.set(qn('w:w'),'90');edge.set(qn('w:type'),'dxa');margins.append(edge)
                pr.append(margins)
                for pp in cell.paragraphs:
                    pp.paragraph_format.space_after=Pt(3);pp.paragraph_format.line_spacing=1.02
                    if ci>0:pp.alignment=WD_ALIGN_PARAGRAPH.CENTER
                    for r in pp.runs:
                        r.font.size=Pt(10);r.font.bold=ri==0
                        r.font.color.rgb=RGBColor.from_string('FFFFFF' if ri==0 else '000000')
            trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
        tb.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        doc.add_paragraph().paragraph_format.space_after=Pt(0)
    doc.add_paragraph('Predicting World Cup match goal difference',style='Title')
    para('HIT140 Group 81\nObjective 2.1\nBirat Timsina\n13 October 2026')
    heading('Research question')
    para('How well do eight differences in competitive form, performance reliability, past opposition, workload and tournament experience explain and predict the signed goal difference in FIFA World Cup 2026 matches?')
    para('The project compiles 104 official match records and fits an eight-variable linear regression using information from before each match. Every feature compares Team A with Team B. This design expresses a relative advantage and guarantees that reversing the teams reverses the model prediction.')
    heading('Main finding')
    para(f'The model has mixed predictive value. Across the later 22 matches, its mean absolute error is {summary["later"]["MAE"]:.3f} goals, slightly worse than predicting a draw for every match. Its root mean squared error is lower, so it reduces some larger mistakes. Feature-removal tests identify both useful and unsuccessful hypotheses.')
    table(['Later 22 matches','MAE','RMSE','R squared'],[[r.model,f'{r.MAE:.3f}',f'{r.RMSE:.3f}',f'{r.R2:.3f}'] for r in metrics[metrics.split=='later matches'].itertuples()],[3.5,1.0,1.0,1.3])
    para('This is a retrospective redesign: the same later-match block was examined in the original project. The numerical evaluation is reproducible, but it is not an untouched external test. The final overlap count with Objective 2.2 requires the teammate\'s final feature definitions.')

    new('Data sources and preparation')
    para('FIFA\'s public match API supplies match identifiers, team order, dates, stage and final scores. A preserved snapshot contains all 104 matches involving 48 teams. The martj42 international-results repository supplies pre-tournament history. Its scores include extra time and exclude shootouts (FIFA, n.d.; martj42, n.d.).')
    para('The downloaded historical snapshot also contains later results. The code discards all external records from 11 June 2026 onward and uses FIFA for the current tournament. Each feature then applies its own strict earlier-date cutoff. Retrospective downloads reconstruct knowable match information, but do not prove what every website displayed at the original kickoff time.')
    doc.add_heading('Wrangling sequence',2)
    for t in ['Flatten the FIFA JSON and check unique match IDs, 104 rows, 48 teams and complete scores.',
              'Standardise eight team-name aliases, such as Korea Republic to South Korea.',
              'Reconcile all World Cup scores by date and team identity. Three source records reverse home and away order; retain FIFA\'s order.',
              'Create one historical record per team per past match, recording goals for, goals against and synthetic 3/1/0 result points.',
              'Apply an explicit list of major competitive tournaments for recent form. Friendlies and regional invitational cups do not enter the ten-game competitive window.',
              'Compute all eight team components, subtract B from A and join to the target by match_id.']:
        doc.add_paragraph(t,style='List Number')
    doc.add_heading('Target and timing',2)
    para('The target is Team A goals minus Team B goals. Team A is FIFA\'s designated Home entry, not necessarily a genuine home team. Recorded goals include extra time and never include shootout kicks. If the course requires a 90-minute target, regulation-time scores must replace these outcomes and the model must be rerun.')
    para('All historical dates must be strictly earlier than the upcoming match\'s local calendar date. Date-only history means same-day information is conservatively excluded. An opponent\'s earlier form is calculated before the historical meeting with that opponent, preventing a second, less obvious route for future information to enter.')
    para('Every row has the required history; missing values are not replaced with zero. Audit files preserve 2,080 main historical records and 20,800 nested opponent records. Workload and previous World Cup participation have separate event-level audits.')

    for start,title in [(0,'Features describing competitive performance'),(4,'Features describing change and context')]:
        new(title)
        para('For every feature below, the final predictor is the Team A component minus the Team B component. Unless stated otherwise, the history consists of ten earlier competitive internationals.')
        for name,key,construction,reason,limit in FEATURE_INFO[start:start+4]:
            doc.add_heading(name,2)
            para(construction+' '+reason)
            row=ab.loc[key]
            reading='removal increased error' if row.delta_MAE>0 else 'removal reduced error'
            para(limit+f' Validation evidence: removing this feature changes pooled MAE by {row.delta_MAE:+.3f} goals; {reading}.')

    new('Model and evaluation design')
    para('The primary model is ordinary least-squares linear regression with exactly eight inputs. Its prediction is the sum of each feature gap multiplied by its fitted coefficient. The intercept is fixed at zero and scaling divides by training standard deviations without subtracting the training means. Consequently, reversing every gap reverses the prediction. This structural choice avoids a preference for an arbitrary Team A label.')
    para('Scikit-learn fits the scaler only on each training subset. The model uses 82 matches before 2 July 2026 for its fixed training fit and evaluates the later 22 matches. The existing cutoff is retained rather than searching for a more favourable split. No team names, match IDs, dates, current-match statistics or stage indicators enter the eight-column predictor matrix (scikit-learn developers, n.d.).')
    doc.add_heading('Training period validation',2)
    para('Two expanding validation folds assess incremental feature usefulness. The first uses 40 training matches and 26 validation matches; the second uses 66 and 16. All training dates precede their corresponding validation dates. Pooled MAE weights the 42 validation predictions equally, rather than weighting differently sized folds equally.')
    para('Each ablation removes one feature and refits on the same folds. Positive removal delta means that losing the feature increases validation error. Seven-variable diagnostic fits do not change the primary eight-variable model. Sensitivity checks vary history length, include friendlies and allow an intercept. These checks evaluate the assumptions; they do not establish universal optimal settings.')
    doc.add_heading('Sequential backtest',2)
    para('A second evaluation starts after 40 matches and refits before each new calendar date. It produces 64 predictions from earlier information only. Earlier completed evaluation matches can enter later rolling histories. This represents sequential prediction before each kickoff, not forecasting every match before the tournament starts.')
    para(f'Sequential MAE is {summary["sequential"]["MAE"]:.3f}, compared with {summary["sequential_zero"]["MAE"]:.3f} for always predicting zero. This is encouraging within this reconstruction, but feature design occurred after the tournament and earlier project results were known. Both evaluations remain exploratory.')
    doc.add_heading('Meaning of the metrics',2)
    para('MAE is the average absolute margin error, in goals. RMSE gives larger mistakes more weight. Lower values are better for both. R squared compares squared errors with a constant equal to the evaluation outcomes\' mean. Negative values are possible. This mathematical reference differs from the practical baseline based only on the training mean.')

    new('Prediction results')
    para(f'Later-match MAE is {summary["later"]["MAE"]:.3f} versus {summary["zero_baseline"]["MAE"]:.3f} for the draw baseline. RMSE is {summary["later"]["RMSE"]:.3f} versus {summary["zero_baseline"]["RMSE"]:.3f}. This supports a limited, metric-dependent finding rather than a general claim of high accuracy.')
    doc.add_picture(str(ROOT/'charts/baselines.png'),width=Inches(6.65))
    para('Figure 1. Errors on the later 22 matches. All values come from the saved predictions.')
    doc.add_picture(str(ROOT/'charts/diagnostics.png'),width=Inches(6.65))
    para('Figure 2. Prediction and residual plots. The small evaluation sample limits conclusions about residual patterns. Errors are actual minus predicted goal difference.')

    new('Systematic feature evaluation')
    para('Scoring-failure rate improved both validation folds. World Cup experience improved pooled MAE but had opposite effects across folds, showing that its apparent value is unstable. The remaining six features failed to improve pooled MAE. These are findings to explain, not values to hide.')
    table(['Removed feature','Pooled delta','Fold 1 delta','Fold 2 delta'],
          [[r.label,f'{r.delta_MAE:+.3f}',f'{r.fold_1_delta:+.3f}',f'{r.fold_2_delta:+.3f}'] for r in ab.reset_index().itertuples()],
          [3.5,1.1,1.1,1.1])
    para('Positive delta means removal worsens MAE. Negative delta means removal improves MAE. The two folds provide limited evidence and correlated features can substitute for each other.')
    table(['Sensitivity variant','Validation MAE'],[[r.variant,f'{r.cv_MAE:.3f}'] for r in sensitivity.itertuples()],[5.4,1.4])
    para('Ten games performed better than six or fourteen under the competitive design, but allowing all internationals performed better still. The competitive-only choice improves consistency of match incentives but sacrifices useful observations in this evaluation. An intercept also improved validation MAE, at the cost of exact reversal symmetry. The primary model retains its stated structural assumptions and reports these trade-offs.')
    para(f'Training VIF values range from {coeff.VIF.min():.2f} to {coeff.VIF.max():.2f}. These are centred correlation diagnostics, not proof that coefficients are stable. Competitive form and clean sheets share information. A negative clean-sheet coefficient does not show that keeping a clean sheet causes poor performance.')

    new('Limitations and assignment requirements')
    para('This design is distinct from the teammate\'s reported xG and Elo approach, but it still derives several measures from historical results. Shared source data alone does not make predictors identical. Conversely, a new name or difference transformation does not establish uniqueness. The exact eight 2.2 definitions and the lecturer\'s interpretation are needed to certify the maximum-four rule.')
    table(['Requirement','Evidence'],[['104 matches','104 unique official IDs and score reconciliation'],['Exactly eight predictors','Explicit feature list and ten-column model CSV'],['Information before the match','Exclusive date filters and historical audits'],['Linear regression and evaluation','OLS, temporal folds, baselines and ablations'],['At most four shared with 2.2','Conditional on final teammate definitions']],[2.6,4.2])
    para('Historical results are noisy indicators of strength. Heuer and Rubner (2014), studying German league football, found chances more informative than goals in their setting. That motivates treating these outcome-based features as hypotheses. It does not prove that our eight engineered variables are superior to xG.')
    para('Repeated teams and overlapping windows create dependent observations. Late knockout matches also differ from the largely group-stage training period. With 104 matches and only two validation folds, small metric changes should not be presented as decisive evidence or causal effects. No coefficient significance claims are made.')
    para('Schedule strength remains confederation-dependent, and there is no adjustment for player injuries, line-ups or individual minutes. Tournament experience partly reflects national strength and does not represent current squad experience. Volatility may be better suited to modelling uncertainty than the mean margin.')
    para('The selected competitive-event list excludes some regional competitions. Workload includes all recorded internationals but omits club matches. Fixture-congestion research has mixed findings, so the workload feature is a testable proxy rather than a guaranteed fatigue effect (Julian et al., 2021).')
    para('Before submission, confirm the external historical source is permitted and the extra-time target convention is appropriate. Record the final overlap comparison in the course Teams space. Review the AI-assisted code and complete the official AI usage declaration with actual contributions and conversation links.')

    new('Reproduction and references')
    para('Run project.py to rebuild the dataset, evaluations and charts. Run test_project.py to verify 12 integrity properties, including future-result invariance and team-reversal symmetry. Then run build_deliverables.py to refresh the report and notebook. Requirements and the actual software versions are included. The executed notebook provides a step-by-step explanation and saved outputs.')
    para('The raw snapshots and their original source manifest are preserved from the earlier acquisition. Checksums identify the exact bytes. Downloading newer files can change the results; use download_sources.py to save a separate new snapshot instead of silently overwriting the original.')
    doc.add_heading('References',2)
    refs=[
      'FIFA. (n.d.). FIFA World Cup 2026 match calendar [Data set]. https://api.fifa.com/api/v3/calendar/matches?idSeason=285023&count=200&language=en',
      'Heuer, A., & Rubner, O. (2014). Optimizing the prediction process: From statistical concepts to the case study of soccer. PLOS ONE, 9(9), e104647. https://doi.org/10.1371/journal.pone.0104647',
      'Julian, R., Page, R. M., & Harper, L. D. (2021). The effect of fixture congestion on performance during professional male soccer match-play: A systematic critical review with meta-analysis. Sports Medicine, 51, 255–273. https://doi.org/10.1007/s40279-020-01359-9',
      'martj42. (n.d.). International football results [Data set]. GitHub. https://github.com/martj42/international_results',
      'scikit-learn developers. (n.d.). LinearRegression. https://scikit-learn.org/stable/modules/generated/sklearn.linear_model.LinearRegression.html',
      'scikit-learn developers. (n.d.). Common pitfalls and recommended practices. https://scikit-learn.org/stable/common_pitfalls.html',
    ]
    for ref in refs:para(ref)
    para('The research motivates broad considerations and limitations. The eight exact formulas, competition filter and timing rules are project design choices evaluated here; the cited studies do not validate this particular feature set.')
    doc.core_properties.author='Birat Timsina'
    doc.core_properties.title='Predicting World Cup match goal difference'
    doc.save(OUT/'Objective_2_1_Report.docx')
    dictionary=pd.DataFrame(FEATURE_INFO,columns=['name','feature','construction','reason','limitation'])
    dictionary.to_csv(ROOT/'data/feature_dictionary.csv',index=False)
    slide_data={'summary':summary,'metrics':metrics.to_dict('records'),'ablation':ab.reset_index().to_dict('records'),
                'features':dictionary.to_dict('records')}
    (ROOT/'results/slide_data.json').write_text(json.dumps(slide_data,indent=2))

def create_notebook():
    cells=[]
    def md(s):cells.append({'cell_type':'markdown','metadata':{},'source':s})
    def code(s):cells.append({'cell_type':'code','metadata':{},'source':s,'outputs':[],'execution_count':None})
    md('# Group 81 Objective 2.1\nBirat Timsina. Eight differences describing competitive form and match context. This executed notebook uses the same project.py functions as the scripts. Read REQUIREMENTS_AND_OVERLAP.md before submission.')
    code('import project as p\nimport pandas as pd\nmatches, history = p.load_data()\nprint("Official matches:", len(matches))\ndisplay(matches.head())')
    md('## Build the eight features\nAll date comparisons use strictly earlier dates. The target is A goals minus B goals, including extra time and excluding shootouts. The model file has one ID, eight X columns and y.')
    code('data = p.build_features(matches, history)\nprint("Dataset shape:", data.shape)\ndisplay(data.head())\ndisplay(pd.read_csv(p.ROOT / "data/feature_dictionary.csv")[["name", "construction"]])')
    md('## Trace a worked example\nThe component file contains each team\'s value before subtraction. The main audit lists the ten competitive games. The separate opponent audit lists the ten games used for each historical opponent.')
    code('first_id = matches.match_id.iloc[0]\ncomponents = pd.read_csv(p.ROOT / "data/team_components_208.csv", dtype={"match_id": str})\naudit = pd.read_csv(p.ROOT / "data/rolling_history_2080.csv", dtype={"match_id": str})\ndisplay(components[components.match_id == first_id])\ndisplay(audit[audit.match_id == first_id][["side", "date", "team", "opponent", "gf", "ga", "points", "opponent_prior_ppg"]])')
    md('## Fit the model\nScaling is learned from the training subset only. No intercept and no centering ensure that swapping teams reverses the prediction. The 22 later games were previously inspected in version 1, so this is a retrospective evaluation.')
    code('train, later, folds = p.splits(matches)\nmodel = p.model().fit(data.loc[train, p.FEATURES], data.loc[train, p.TARGET])\nprediction = model.predict(data.loc[later, p.FEATURES])\nprint(p.scores(data.loc[later, p.TARGET], prediction))\ndisplay(pd.read_csv(p.ROOT / "results/metrics.csv"))')
    md('## Evaluate each feature\nPositive removal delta means losing the feature worsens validation MAE. Read both fold deltas because pooled improvements can hide instability.')
    code('display(pd.read_csv(p.ROOT / "results/feature_ablation.csv"))\ndisplay(pd.read_csv(p.ROOT / "results/sensitivity.csv"))')
    md('![Feature removal results](charts/ablation.png)')
    md('## Inspect prediction errors\nThe model is slightly worse on MAE and better on RMSE than the zero baseline. Do not report this as high accuracy.')
    code('display(pd.read_csv(p.ROOT / "results/sequential_metrics.csv"))')
    md('![Baseline comparison](charts/baselines.png)\n\n![Prediction diagnostics](charts/diagnostics.png)')
    md('## Verify integrity\nThese tests check data shape and temporal logic. They do not certify the correctness of every historical source result or unseen course requirements.')
    code('import unittest\nimport test_project\nsuite = unittest.defaultTestLoader.loadTestsFromModule(test_project)\ncheck = unittest.TextTestRunner(verbosity=1).run(suite)\nassert check.wasSuccessful()')
    # Execute the exact Python cells in a shared namespace, capturing real output.
    # The host lacks Jupyter; the display hook serialises actual DataFrame values.
    namespace={};number=0
    for i,cell in enumerate(cells):
        cell['id']=f'objective21-{i}'
        if cell['cell_type']!='code':continue
        number+=1;cell['execution_count']=number;outputs=[]
        def display(value):
            data={'text/plain':str(value)}
            if hasattr(value,'to_html'):data['text/html']=value.to_html(index=False)
            outputs.append({'output_type':'display_data','data':data,'metadata':{}})
        namespace['display']=display
        stdout=io.StringIO();stderr=io.StringIO()
        with contextlib.redirect_stdout(stdout),contextlib.redirect_stderr(stderr):
            exec(compile(cell['source'],f'notebook_cell_{i}','exec'),namespace)
        if stdout.getvalue():outputs.append({'output_type':'stream','name':'stdout','text':stdout.getvalue()})
        if stderr.getvalue():outputs.append({'output_type':'stream','name':'stderr','text':stderr.getvalue()})
        cell['outputs']=outputs
    nb={'cells':cells,'nbformat':4,'nbformat_minor':5,
        'metadata':{'kernelspec':{'name':'python3','display_name':'Python 3','language':'python'},
                    'execution_note':'Exact cells executed in a shared Python namespace; display outputs captured by the builder.'}}
    (ROOT/'WALKTHROUGH.ipynb').write_text(json.dumps(nb,indent=2))

if __name__=='__main__':
    create_report();create_notebook()
    result=subprocess.run([sys.executable,str(ROOT/'test_project.py')],cwd=ROOT,capture_output=True,text=True)
    (ROOT/'results/test_log.txt').write_text(result.stdout+result.stderr)
    assert result.returncode==0
    print('Report, executed notebook, slide data and test log created.')
